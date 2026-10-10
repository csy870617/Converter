# PDF → Word(DOCX) 변환. 브라우저(Pyodide)에서 pdf2docx를 돌리고, 결과 품질을 떨어뜨리는 문제를 고친다.
#
# pdf2docx(MIT)는 PDF의 글자·표·그림·단 나누기를 Word 문단·표·그림으로 다시 짜는 도구다.
# 그대로 쓰면 다음 문제가 있어 여기서 덧씌운다:
#  1) 단어 사이 띄어쓰기가 사라진다: 새 PyMuPDF는 글자 사이 간격으로 추정한 띄어쓰기를 별도 조각으로
#     주는데, pdf2docx가 공백뿐인 조각을 버린다.
#  2) 그림(벡터)이 많은 페이지가 수십 초씩 걸린다: 그림 개수만큼 페이지 전체 그림 목록을 다시 훑는다.
#  3) 글꼴 이름이 PDF 안 이름(NimbusRomNo9L, ArialMT …)이라 Word에 없는 글꼴이 되어 모양이 바뀐다.
#  4) 글머리표가 Symbol/OpenSymbol 전용 문자(U+F0B7 등)로 남아 네모(□)로 보인다.
#  5) 제목·글머리표 항목이 다음 문단에 붙는다: 문장부호로 끝나야만 문단을 나눈다.

import logging
import re

import pymupdf
import pdf2docx
from docx.oxml.ns import qn
from pdf2docx.common import constants
from pdf2docx.image.ImagesExtractor import ImagesExtractor
from pdf2docx.page.RawPageFitz import RawPageFitz
from pdf2docx.text.Lines import Lines
from pdf2docx.text.TextSpan import TextSpan
from pdf2docx.text.TextBlock import TextBlock
from pdf2docx.table.TableBlock import TableBlock
from pdf2docx.layout.Blocks import Blocks
from pdf2docx.layout.Column import Column
from pdf2docx.layout.Section import Section
from pdf2docx.common.Collection import Collection
from pdf2docx.page.RawPage import RawPage
from pdf2docx.text.Line import Line
from pdf2docx.table.TablesConstructor import TablesConstructor
from pdf2docx.table.TableStructure import TableStructure
from pdf2docx.table.Cell import Cell
from pdf2docx.table.Row import Row
from docx.enum.table import WD_ROW_HEIGHT
from pdf2docx.common.Element import Element
from pdf2docx.common.share import rgb_component
from pdf2docx.page.Page import Page
from pdf2docx.page.Pages import Pages
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import RGBColor
from pdf2docx.shape.Paths import Paths
from pdf2docx.image.ImageSpan import ImageSpan
from pdf2docx.common.docx import set_columns, reset_paragraph_format
from docx.enum.text import WD_BREAK
from docx.shared import Pt

# ---------------------------------------------------------------------------
# 1) 띄어쓰기 · 4) 글머리표 문자
# ---------------------------------------------------------------------------

# Symbol / Wingdings 글꼴의 전용 영역 문자 → 어느 글꼴에나 있는 문자
PUA_BULLETS = {
    '': '•', '': '•', '': '•', '': '▪', '': '◦', '': '■', '': '❑',
    '': '◆', '': '❖', '': '➢', '': '→', '': '➔', '': '✓', '': '✗',
    '': '-', '': '○', '': '·', '': '◇',
    '': '•', '': '▪',  # OpenSymbol
}


_STYLE_KEYS = ('font', 'size', 'flags', 'color', 'alpha')
_SYNTHETIC = 4 | 512  # MuPDF가 간격을 보고 넣은 띄어쓰기 (FZ_STEXT_SYNTHETIC, FZ_STEXT_SYNTHETIC_LARGE)
_HANGUL = re.compile('[\u1100-\u11ff\u3130-\u318f\uac00-\ud7a3]')


def _same_style(a, b):
    return (a.get('font') == b.get('font') and abs(a.get('size', 0) - b.get('size', 0)) < 0.01
            and a.get('flags') == b.get('flags') and a.get('color') == b.get('color')
            and a.get('alpha', 255) == b.get('alpha', 255))


def _union(a, b):
    return (min(a[0], b[0]), min(a[1], b[1]), max(a[2], b[2]), max(a[3], b[3]))


def _horizontal(line):
    d = line.get('dir', (1, 0))
    return line.get('wmode', 0) == 0 and abs(d[0] - 1) < 1e-3 and abs(d[1]) < 1e-3


def _line_size(line):
    return max((sp.get('size', 0) for sp in line.get('spans', [])), default=0) or 1


def _vertical_overlap(a, b):
    h = min(a[3] - a[1], b[3] - b[1])
    return h > 0 and min(a[3], b[3]) - max(a[1], b[1]) > 0.5 * h


def _rebuild(lines):
    """같은 줄인데 따로 그려진 조각들을 글자 위치 순서대로 한 줄로 합친다. 겹쳐 그린 글자(굵게 흉내 등)면 None."""
    chars = [(ch, sp) for line in lines for sp in line['spans'] for ch in sp.get('chars', [])]
    chars.sort(key=lambda t: (t[0]['bbox'][0] + t[0]['bbox'][2]) / 2)
    for (a, _), (b, _) in zip(chars, chars[1:]):
        w = min(a['bbox'][2] - a['bbox'][0], b['bbox'][2] - b['bbox'][0])
        if w > 0 and a['c'] == b['c'] and not a['c'].isspace() and min(a['bbox'][2], b['bbox'][2]) - max(a['bbox'][0], b['bbox'][0]) > 0.5 * w:
            return None
    # 두 번 그려진 띄어쓰기(같은 자리)는 하나만 남긴다
    deduped = []
    for ch, src in chars:
        if deduped and ch['c'].isspace() and deduped[-1][0]['c'].isspace():
            a = deduped[-1][0]['bbox']
            if min(a[2], ch['bbox'][2]) - max(a[0], ch['bbox'][0]) > 0.3 * max(a[2] - a[0], ch['bbox'][2] - ch['bbox'][0], 0.1):
                continue
        deduped.append((ch, src))
    chars = deduped
    spans = []
    for ch, src in chars:
        last = spans[-1] if spans else None
        if last is not None and (last['_src'] is src or _same_style(last, src)):
            last['chars'].append(ch)
            last['bbox'] = _union(last['bbox'], ch['bbox'])
        else:
            spans.append(dict(src, chars=[ch], bbox=tuple(ch['bbox']), origin=ch.get('origin', src.get('origin')), _src=src))
    for sp in spans:
        del sp['_src']
    bbox = lines[0]['bbox']
    for line in lines[1:]:
        bbox = _union(bbox, line['bbox'])
    return dict(lines[0], spans=spans, bbox=bbox)


def merge_row_pieces(block):
    """한 줄을 라틴 글자·한글을 따로따로 그린 PDF(LibreOffice·한글 등)는 MuPDF가 한 줄을 여러 조각으로 나눈다.
    조각마다 띄어쓰기만 든 조각도 생기는데 pdf2docx는 이를 버려 띄어쓰기가 사라지고 글자 순서도 섞인다.
    같은 높이(행)에 있고 서로 겹치거나 가까운 조각을 글자 위치 순서대로 한 줄로 다시 묶는다."""
    lines = block.get('lines', [])
    if len(lines) < 2:
        return
    order = {id(line): i for i, line in enumerate(lines)}
    horiz = [line for line in lines if _horizontal(line) and line.get('spans')]
    rows = []
    for line in sorted(horiz, key=lambda l: (l['bbox'][1] + l['bbox'][3]) / 2):
        row = rows[-1] if rows else None
        if row and _vertical_overlap(row['bbox'], line['bbox']):
            row['lines'].append(line)
            row['bbox'] = _union(row['bbox'], line['bbox'])
        else:
            rows.append({'bbox': line['bbox'], 'lines': [line]})
    groups = []
    for row in rows:
        current, right = [], None
        for line in sorted(row['lines'], key=lambda l: l['bbox'][0]):
            size = _line_size(line)
            if current and line['bbox'][0] - right < 0.6 * max(size, max(_line_size(o) for o in current)):
                current.append(line)
                right = max(right, line['bbox'][2])
            else:
                if current:
                    groups.append(current)
                current, right = [line], line['bbox'][2]
        if current:
            groups.append(current)
    merged = {}
    for g in groups:
        if len(g) > 1:
            line = _rebuild(g)
            if line:
                first = min(g, key=lambda l: order[id(l)])
                merged[id(first)] = line
                for other in g:
                    if other is not first:
                        merged[id(other)] = None
    block['lines'] = [merged.get(id(line), line) for line in lines if merged.get(id(line), line) is not None]


def infer_korean_spaces(line):
    """띄어쓰기 글자 없이 글자 위치만으로 띄어 쓴 한글 PDF(한글 프로그램에서 만든 PDF 등)에 띄어쓰기를 넣는다.
    MuPDF는 한중일 글자 사이에는 띄어쓰기를 추정하지 않는다. 줄 안의 글자 간격을 '단어 안'과 '단어 사이'
    두 무리로 나눠(2-평균), 큰 쪽 간격에만 넣는다. 간격이 고른 줄(자간을 넓힌 제목 등)은 건드리지 않는다."""
    pairs = []
    flat = [(ch, sp) for sp in line.get('spans', []) for ch in sp.get('chars', [])]
    if len(flat) < 3 or not _HANGUL.search(''.join(ch['c'] for ch, _ in flat)):
        return
    if any(ch['c'] == ' ' and not ch.get('_syn') for ch, _ in flat):
        return  # 진짜 띄어쓰기 글자가 있는 PDF: 그대로 믿는다
    for (a, sa), (b, sb) in zip(flat, flat[1:]):
        if a['c'].isspace() or b['c'].isspace():
            continue
        # 한글 옆 간격만 본다 (영문끼리의 띄어쓰기는 MuPDF가 넣고, 글꼴이 바뀐 영문은 간격이 들쭉날쭉해 기준을 흐린다)
        if not (_HANGUL.match(a['c']) or _HANGUL.match(b['c'])):
            continue
        size = max(sa.get('size', 0), sb.get('size', 0)) or 1
        pairs.append(((b['bbox'][0] - a['bbox'][2]) / size, a, sa))
    if len(pairs) < 2:
        return
    vals = sorted(v for v, _, _ in pairs)
    lo, hi = vals[0], vals[-1]
    if hi - lo < 0.15:
        return
    c1, c2 = lo, hi
    for _ in range(20):
        t = (c1 + c2) / 2
        low = [v for v in vals if v <= t]
        high = [v for v in vals if v > t]
        if not low or not high:
            return
        c1, c2 = sum(low) / len(low), sum(high) / len(high)
    if c2 - c1 < 0.15:
        return
    t = max((c1 + c2) / 2, c1 + 0.1)
    for gap, a, sa in pairs:
        if gap > t:
            i = next(k for k, ch in enumerate(sa['chars']) if ch is a)
            x0, y0, x1, y1 = a['bbox']
            space = dict(a, c=' ', bbox=(x1, y0, x1 + gap * (sa.get('size', 0) or 1), y1), _syn=True)
            sa['chars'].insert(i + 1, space)


_CJK = re.compile('[\u1100-\u11ff\u3040-\u30ff\u3130-\u318f\u3400-\u9fff\uac00-\ud7a3\uf900-\ufaff]')
_ALNUM = re.compile(r'[A-Za-z0-9]')


def drop_autospace(line):
    """Word·LibreOffice는 한글과 숫자·영문 사이를 조금 띄워 그린다(자동 간격). MuPDF는 이를 띄어쓰기로 추정해
    '2026년'이 '2026 년'이 된다. 진짜 띄어쓰기 글자를 쓰는 PDF(그 줄에 진짜 띄어쓰기가 있음)라면,
    한글↔숫자·영문 경계의 좁은 추정 띄어쓰기는 지운다. (넓은 간격은 탭일 수 있어 남긴다)"""
    flat = [(ch, sp) for sp in line.get('spans', []) for ch in sp.get('chars', [])]
    if not any(ch['c'] == ' ' and not ch.get('_syn') for ch, _ in flat):
        return
    drop = set()
    for i in range(1, len(flat) - 1):
        ch, sp = flat[i]
        if not (ch.get('_syn') and ch['c'] == ' '):
            continue
        a, b = flat[i - 1][0], flat[i + 1][0]
        if a['c'].isspace() or b['c'].isspace():
            continue
        size = sp.get('size', 0) or 1
        # 자동 간격은 한글과 영문·숫자 사이에만 생긴다 (괄호·쉼표 같은 문장 부호 옆의 띄어쓰기는 진짜다)
        pair = (a['c'], b['c']) if _CJK.match(a['c']) else (b['c'], a['c'])
        if (b['bbox'][0] - a['bbox'][2]) / size < 0.6 and _CJK.match(pair[0]) and _ALNUM.match(pair[1]):
            drop.add(id(ch))
    if drop:
        for sp in line.get('spans', []):
            sp['chars'] = [ch for ch in sp.get('chars', []) if id(ch) not in drop]


_MIDDLE_DOTS = '·‧∙・ㆍ'


def drop_dot_spaces(line):
    """가운뎃점 옆의 추정 띄어쓰기를 지운다. 점 글자가 자리보다 좁으면 옆이 비어 MuPDF가 띄어쓰기로 보고
    '도·소매업'이 '도· 소매업'이 된다. 낱말 사이의 점일 때만 (줄 첫머리의 '· 항목' 같은 글머리표는 그대로)."""
    flat = [(ch, sp) for sp in line.get('spans', []) for ch in sp.get('chars', [])]
    drop = set()
    for i in range(1, len(flat) - 1):
        ch, sp = flat[i]
        if not (ch.get('_syn') and ch['c'] == ' '):
            continue
        a, b = flat[i - 1][0], flat[i + 1][0]
        if not a['c'].strip() or not b['c'].strip():
            continue
        if a['c'] in _MIDDLE_DOTS:
            word = i >= 2 and flat[i - 2][0]['c'].strip()
        elif b['c'] in _MIDDLE_DOTS:
            word = i + 2 < len(flat) and flat[i + 2][0]['c'].strip()
        else:
            continue
        size = sp.get('size', 0) or 1
        if word and (b['bbox'][0] - a['bbox'][2]) / size < 1.0:
            drop.add(id(ch))
    if drop:
        for sp in line.get('spans', []):
            sp['chars'] = [ch for ch in sp.get('chars', []) if id(ch) not in drop]


def tidy_text_blocks(blocks):
    """MuPDF가 준 글자 정보를 pdf2docx가 다루기 전에 바로잡는다."""
    for block in blocks:
        for line in block.get('lines', []):
            for span in line.get('spans', []):
                syn = bool(span.get('char_flags', 0) & _SYNTHETIC)
                for ch in span.get('chars', []):
                    if ch.get('c') in PUA_BULLETS:
                        ch['c'] = PUA_BULLETS[ch['c']]
                    # 추정 띄어쓰기가 다음 글자 조각에 섞여 들어가면 조각 표시로는 모른다: 글자마다 표시도 본다
                    if syn or ch.get('synthetic'):
                        ch['_syn'] = True
        merge_row_pieces(block)
        for line in block.get('lines', []):
            drop_autospace(line)
            infer_korean_spaces(line)
            drop_dot_spaces(line)
            horizontal = line.get('dir', (1, 0))[0] > 0.99
            merged = []
            for span in line.get('spans', []):
                chars = span.get('chars', [])
                blank = not ''.join(ch.get('c', '') for ch in chars).strip()
                if not blank and horizontal:
                    spacing = char_spacing(span)
                    if spacing:
                        span['char_spacing'] = spacing
                # 자간이 다른 조각은 합치지 않는다 (표의 칸마다 자간이 달라 합치면 평균이 되어 칸을 넘친다)
                if merged and (blank or (_same_style(merged[-1], span)
                                         and abs(merged[-1].get('char_spacing', 0) - span.get('char_spacing', 0)) < 0.1)):
                    prev = merged[-1]
                    prev['chars'] = prev.get('chars', []) + chars
                    prev['bbox'] = _union(prev['bbox'], span['bbox'])
                    continue
                merged.append(span)
            line['spans'] = merged
            for span in merged:
                for ch in span.get('chars', []):
                    ch.pop('_syn', None)
    return blocks


def char_spacing(span):
    """자간(글자 사이 간격, pt). 한컴 오피스 문서처럼 글자 사이를 좁히거나 넓힌 글은 Word에서도 같은 폭이 되게 한다.
    (그대로 두면 Word에서 글이 길어져 좁은 표 칸에서 줄이 넘어간다.) 낱말 사이 빈칸과 칸 건너뛰기는 세지 않는다."""
    chars = span.get('chars', [])
    limit = 0.5 * span.get('size', 10)
    gaps = []
    for a, b in zip(chars, chars[1:]):
        if not a.get('c', '').strip() or not b.get('c', '').strip():
            continue
        gap = b['origin'][0] - a['origin'][0] - (a['bbox'][2] - a['bbox'][0])
        if abs(gap) <= limit:
            gaps.append(gap)
    if not gaps:
        return 0.0
    mean = sum(gaps) / len(gaps)
    if abs(mean) < 0.15:
        return 0.0
    # Word는 마지막 글자 뒤에도 자간을 더하므로 글자 수에 맞춰 나누고, 2% 여유를 둔다
    # (글꼴이 조금만 넓어도 좁은 칸에서 끝 글자가 다음 줄로 넘어간다)
    widths = [c['bbox'][2] - c['bbox'][0] for c in chars if c.get('c', '').strip()]
    spacing = mean * len(gaps) / (len(gaps) + 1) - 0.02 * sum(widths) / len(widths)
    if mean > 0:
        spacing = max(0.0, spacing)
    return round(spacing, 2) if abs(spacing) >= 0.1 else 0.0


_preprocess_text = RawPageFitz._preprocess_text


def _preprocess_text_patched(self, **settings):
    return drop_margin_stamps(tidy_text_blocks(_preprocess_text(self, **settings)))


RawPageFitz._preprocess_text = _preprocess_text_patched

# ---------------------------------------------------------------------------
# 2) 그림 많은 페이지 속도
# ---------------------------------------------------------------------------


def _hide_page_text_and_images(page, rm_text, rm_image):
    """원본과 같지만, 페이지의 그림 이름 목록을 한 번만 구한다 (원본은 개체마다 다시 구해 O(n²))."""
    xref_list = [xref for (xref, name, invoker, bbox) in page.get_xobjects()]
    xref_list.extend(page.get_contents())
    img_names = [item[7] for item in page.get_images(full=True)] if rm_image else []

    def hide_text(stream):
        res = stream
        found = False
        for k in ['BT', 'Tm', 'Td', '2 Tr']:
            bk = k.encode()
            if bk in stream:
                found = True
                res = res.replace(bk, f'{k} 3 Tr'.encode())
        return res, found

    def hide_images(stream):
        res = stream
        found = False
        for k in img_names:
            bk = f'/{k} Do'.encode()
            if bk in stream:
                found = True
                res = res.replace(bk, b'')
        return res, found

    doc = page.parent
    source = {}
    for xref in xref_list:
        src = doc.xref_stream(xref)
        stream, found_text = hide_text(src) if rm_text else (src, False)
        stream, found_images = hide_images(stream) if rm_image else (stream, False)
        if found_text or found_images:
            doc.update_stream(xref, stream)
            source[xref] = src  # 그린 뒤 원래대로 되돌릴 수 있게 원본을 돌려준다
    return source


ImagesExtractor._hide_page_text_and_images = staticmethod(_hide_page_text_and_images)

# ---------------------------------------------------------------------------
# 3) 글꼴 이름 → Word에서 쓰는 이름
# ---------------------------------------------------------------------------

# (정규식, Word 글꼴). 위에서부터 먼저 맞는 것을 쓴다. 대소문자·공백·하이픈은 무시한다.
FONT_RULES = [
    # 한글 글꼴 (Windows·한컴 오피스 이름)
    (r'malgun|맑은고딕', '맑은 고딕'),
    (r'batangche|바탕체', '바탕체'), (r'batang|바탕', '바탕'),
    (r'gulimche|굴림체', '굴림체'), (r'gulim|굴림', '굴림'),
    (r'dotumche|돋움체', '돋움체'), (r'dotum|돋움', '돋움'),
    (r'gungsuh|궁서', '궁서'),
    (r'hcrbatang|함초롬바탕', '함초롬바탕'), (r'hcrdotum|함초롬돋움', '함초롬돋움'),
    (r'nanumbarungothic|나눔바른고딕', '나눔바른고딕'), (r'nanumsquare|나눔스퀘어', '나눔스퀘어'),
    (r'nanummyeongjo|나눔명조', '나눔명조'), (r'nanumgothiccoding', '나눔고딕코딩'), (r'nanumgothic|나눔고딕', '나눔고딕'),
    (r'applesdgothic|applegothic', '맑은 고딕'), (r'applemyungjo', '바탕'),
    # Adobe 한글 표준 글꼴 이름 (글꼴을 넣지 않은 옛 PDF)
    (r'hygothic|hygoth|hy중고딕|hygungso', '맑은 고딕'), (r'hysmyeongjo|hymyeongjo|hy신명조', '바탕'),
    (r'notoserifcjk|notoserifkr|sourcehanserif', '바탕'),
    (r'notosanscjk|notosanskr|sourcehansans|wenquanyi|droidsansfallback|unifont', '맑은 고딕'),
    # 서양 글꼴
    (r'arialnarrow', 'Arial Narrow'), (r'arialblack', 'Arial Black'),
    (r'arial|helvetica|nimbussan|liberationsans|arimo|texgyreheros|freesans', 'Arial'),
    (r'timesnewroman|times|nimbusrom|liberationserif|tinos|texgyretermes|freeserif|stix', 'Times New Roman'),
    (r'couriernew|courier|nimbusmono|liberationmono|cousine|texgyrecursor|freemono', 'Courier New'),
    (r'carlito|calibri', 'Calibri'), (r'caladea|cambriamath', 'Cambria Math'), (r'cambria', 'Cambria'),
    (r'dejavusansmono|bitstreamverasansmono|menlo|monaco|consolas|sfmono|cmtt|sftt|lmmono|inconsolata', 'Consolas'),
    (r'dejavusans|bitstreamverasans|verdana', 'Verdana'),
    (r'dejavuserif|bitstreamveraserif|georgia', 'Georgia'),
    (r'segoeui', 'Segoe UI'), (r'tahoma', 'Tahoma'), (r'trebuchet', 'Trebuchet MS'),
    (r'garamond', 'Garamond'), (r'bookantiqua|palatino|texgyrepagella|urwpalladio', 'Book Antiqua'),
    (r'centurygothic', 'Century Gothic'), (r'century', 'Century'),
    (r'opensans|roboto|lato|sourcesans|notosans', 'Arial'),
    (r'notoserif|sourceserif|merriweather', 'Times New Roman'),
    # TeX 글꼴 (논문)
    (r'^(cm|sf|lm)(r|bx|b|ti|sl|ss|csc|ssbx|ssi|rm|bxti|u)\d*|lmroman|latinmodern|cmunrm', 'Times New Roman'),
    (r'^(cm|lm)(mi|mib|sy|bsy|ex)\d*|msam|msbm|eufm|rsfs|lmmath|latinmodernmath', 'Cambria Math'),
    (r'symbol', 'Symbol'), (r'zapfdingbats|dingbats|wingdings', 'Wingdings'),
]
_FONT_RULES = [(re.compile(p), name) for p, name in FONT_RULES]
_HANGUL = re.compile('[ᄀ-ᇿ㄰-㆏가-힣]')
_SERIF_HINT = re.compile(r'myeongjo|myungjo|batang|serif|song|ming|mincho|명조|바탕')


def office_font(name, text=''):
    """PDF 안 글꼴 이름을 Word가 아는 글꼴 이름으로 바꾼다. 모르는 글꼴은 이름만 다듬는다."""
    raw = (name or '').split('+')[-1]
    key = re.sub(r'[\s_\-,]', '', raw).lower()
    for pattern, office in _FONT_RULES:
        if pattern.search(key):
            return office
    # 모르는 글꼴이 한글을 담고 있으면 한글 기본 글꼴로 (Word에 없는 글꼴 대신 엉뚱한 글꼴이 쓰이지 않게)
    if _HANGUL.search(text or ''):
        return '바탕' if _SERIF_HINT.search(key) else '맑은 고딕'
    # 굵게·기울임 등 꼬리말 정리 (예: Foo-BoldItalicMT → Foo)
    clean = re.sub(r'[-,](bold|italic|oblique|regular|roman|medium|light|semibold|black|book|regu|medi|ital|bolditalic)\w*$',
                   '', raw, flags=re.I)
    clean = re.sub(r'(PSMT|PS|MT)$', '', clean)
    return clean or raw


_set_text_format = TextSpan._set_text_format


def _set_text_format_patched(self, docx_run):
    original = self.font
    font = office_font(original, self.text)
    self.font = font
    try:
        _set_text_format(self, docx_run)
    finally:
        self.font = original
    # 한글은 동아시아 글꼴 칸이 실제로 쓰인다. 서양 글꼴만 정해져 있으면 한글 기본 글꼴을 함께 적어 둔다.
    if _HANGUL.search(self.text or ''):
        east = font if font in KOREAN_FONTS else ('바탕' if _SERIF_HINT.search(original.lower()) else '맑은 고딕')
        docx_run._element.rPr.rFonts.set(qn('w:eastAsia'), east)
        docx_run._element.rPr.rFonts.set(qn('w:hint'), 'eastAsia')


KOREAN_FONTS = {'맑은 고딕', '바탕', '바탕체', '굴림', '굴림체', '돋움', '돋움체', '궁서', '함초롬바탕', '함초롬돋움',
                '나눔고딕', '나눔명조', '나눔바른고딕', '나눔스퀘어', '나눔고딕코딩'}
TextSpan._set_text_format = _set_text_format_patched

# ---------------------------------------------------------------------------
# 5) 문단 나누기: 제목·글머리표 항목·글자 크기가 바뀌는 곳에서도 나눈다
# ---------------------------------------------------------------------------

_LIST_START = re.compile(
    r'^\s*('
    r'[•●○◦▪▫■□◆◇❖➢➤►▶▷✓✔·∙‣⁃]\s?'          # 글머리표
    r'|[\-–—*]\s'                            # - 항목 (뒤에 빈칸이 있어야: -5% 같은 수는 제외)
    r'|\(?\d{1,3}[.)]\s'                      # 1. 1) (1)
    r'|\(?[a-zA-Z][.)]\s'                     # a. a) (a)
    r'|\(?[가-힣][.)]\s'                      # 가. 가) (가)
    r'|[①-⑳㉠-㉭㈀-㈍]'                       # ① ㉠ ㈀
    r'|[ⅰ-ⅻⅠ-Ⅻ][.)]?\s'                       # 로마 숫자
    r')')


def _text_spans(row):
    return [span for line in row for span in line.spans if isinstance(span, TextSpan) and span.text.strip()]


def _row_size(row):
    sizes = [span.size for span in _text_spans(row)]
    return max(sizes) if sizes else 0


def _row_bold(row):
    spans = _text_spans(row)
    return bool(spans) and all(span.flags & 16 for span in spans)


def _row_text(row):
    return ' '.join(line.text for line in row)


def split_vertically_by_text(self, line_break_free_space_ratio, new_paragraph_free_space_ratio):
    rows = self.group_by_physical_rows()
    for row in rows:
        row.sort_in_line_order()
    num = len(rows)
    if num == 1:
        return rows

    W = max(row[-1].bbox[2] - row[0].bbox[0] for row in rows[1:])
    H = sum(row[0].bbox[3] - row[0].bbox[1] for row in rows) / num
    punc = tuple(constants.SENTENCE_END_PUNC)

    res = []
    lines = Lines()
    prev = None
    for row in rows:
        text = _row_text(row).strip()
        w = row[-1].bbox[2] - row[0].bbox[0]
        start = False
        if prev is not None:
            prev_text = _row_text(prev).strip()
            prev_w = prev[-1].bbox[2] - prev[0].bbox[0]
            size, prev_size = _row_size(row), _row_size(prev)
            prev_short = prev_w / W <= 1.0 - line_break_free_space_ratio
            if prev_text.endswith(punc) and prev_short:
                start = True  # 앞 줄이 문장 끝 + 오른쪽이 비어 있음 → 문단 끝 (원래 규칙)
            elif prev_text.endswith(punc) and (W - w) / H >= new_paragraph_free_space_ratio:
                start = True  # 문장 끝 다음에 짧은(들여쓴) 줄 → 새 문단 (원래 규칙)
            elif _LIST_START.match(text):
                start = True  # 글머리표·번호로 시작하는 줄
            elif size and prev_size and max(size, prev_size) / min(size, prev_size) > 1.15:
                start = True  # 글자 크기가 바뀜 (제목 ↔ 본문)
            elif _row_bold(prev) != _row_bold(row) and prev_short:
                start = True  # 굵은 짧은 줄(소제목) 다음 본문
        if start and lines:
            res.append(lines)
            lines = Lines()
        lines.extend(row)
        prev = row
    if lines:
        res.append(lines)
    return res


Lines.split_vertically_by_text = split_vertically_by_text


def _join_lines_vertically(self, max_line_spacing_ratio):
    """원본은 페이지 전체에서 가장 흔한 줄 간격 하나로만 같은 문단인지 판단해, 그 값이 문단마다 다르면
    한 줄씩 다른 문단이 된다. 줄 높이의 절반보다 가까운 줄은 같은 문단으로 본다."""
    idx0, idx1 = (1, 3) if self.is_horizontal_text else (0, 2)

    def v_bdy(block):
        return block.bbox[idx0], block.bbox[idx1]

    def distance(b1, b2):
        return round(v_bdy(b2)[0] - v_bdy(b1)[1], 2)

    def line_height(line):
        span = max(line.spans, key=lambda s: len(getattr(s, 'text', '') or ''))
        return round(span.bbox.height if line.is_horizontal_text else span.bbox.width, 2)

    def common_spacing():
        if not self._instances:
            return 0.0
        ref1 = v_bdy(self._instances[0])[1]
        distances = []
        for block in self._instances[1:]:
            y0, y1 = v_bdy(block)
            distances.append(round(y0 - ref1, 2))
            ref1 = y1
        return max(distances, key=distances.count) if distances else 0.0

    blocks, lines = [], []

    def close():
        if lines:
            block = TextBlock()
            block.add(lines)
            blocks.append(block)
            lines.clear()

    ref_dis = common_spacing()
    for block in self._instances:
        if isinstance(block, TableBlock):
            close()
            blocks.append(block)
            continue
        ref = lines[-1] if lines else None
        if not ref or ref.in_same_row(block):
            new = False
        elif block.image_spans or ref.image_spans:
            new = True
        else:
            gap, h = distance(ref, block), line_height(ref)
            usual = gap <= ref_dis + 1.0 and ref_dis <= max_line_spacing_ratio * h
            new = not (usual or -0.5 * h <= gap <= 0.5 * h)
        if new:
            close()
        lines.append(block)
    close()
    return blocks


Blocks._join_lines_vertically = _join_lines_vertically

# ---------------------------------------------------------------------------
# 10) 테두리 없는 표 추정: 글줄을 칸 경계에서 자르는 표는 버린다
# ---------------------------------------------------------------------------
# 줄 맞춰 놓은 글(각주의 긴 주소 등)을 표로 추정하면 한 줄이 여러 칸에 걸쳐 잘리고, 칸 밖으로 나간 글자는
# 사라진다(예: 'http://image-net.org/… and' → 'htte2015.'). 그런 표는 만들지 않고 글로 둔다.


def _straddles(table, lines):
    for line in lines:
        lb = line.bbox
        for row in table:
            hits = 0
            for cell in row:
                if cell is None or not getattr(cell, 'bbox', None):
                    continue
                cb = cell.bbox
                if min(lb.y1, cb.y1) - max(lb.y0, cb.y0) <= 0.5 * lb.height:
                    continue
                if min(lb.x1, cb.x1) - max(lb.x0, cb.x0) > 1.0:
                    hits += 1
            if hits > 1:
                return True
    return False


def _stream_tables(self, min_border_clearance, max_border_width, line_separate_threshold):
    table_strokes = self._shapes.table_strokes
    table_fillings = self._shapes.table_fillings
    tables_lines = self._blocks.collect_stream_lines(table_fillings, line_separate_threshold)
    X0, Y0, X1, Y1 = self._parent.bbox

    def top_bottom_boundaries(y0, y1):
        y_lower, y_upper = Y0, Y1
        for block in self._blocks:
            if block.bbox.y1 < y0:
                y_lower = block.bbox.y1
            if block.bbox.y0 > y1:
                y_upper = block.bbox.y0
                break
        return y_lower, y_upper

    tables = Blocks()
    settings = {'min_border_clearance': min_border_clearance, 'max_border_width': max_border_width}
    for table_lines in tables_lines:
        if not table_lines:
            continue
        x0 = min(r.bbox.x0 for r in table_lines)
        y0 = min(r.bbox.y0 for r in table_lines)
        x1 = max(r.bbox.x1 for r in table_lines)
        y1 = max(r.bbox.y1 for r in table_lines)
        y0_margin, y1_margin = top_bottom_boundaries(y0, y1)
        rect = Element().update_bbox((X0, y0_margin, X1, y1_margin))
        explicit_strokes = table_strokes.contained_in_bbox(rect.bbox)
        explicit_shadings, _ = table_fillings.split_with_intersection(rect.bbox, threshold=constants.FACTOR_A_FEW)
        # 표 바깥 테두리는 글자(와 표의 가로줄) 범위에 붙인다. 원본은 단 너비 끝까지 늘려 첫 칸이 좁아진다.
        sx0 = min([x0] + [st.bbox.x0 for st in explicit_strokes])
        sx1 = max([x1] + [st.bbox.x1 for st in explicit_strokes])
        outer_borders = TablesConstructor._outer_borders(
            (x0, y0, x1, y1), (max(X0, sx0 - 3), y0_margin, min(X1, sx1 + 3), y1_margin))
        if not (explicit_shadings or explicit_strokes) and TablesConstructor._is_simple_structure(table_lines):
            continue
        strokes = self._stream_strokes(table_lines, outer_borders, explicit_strokes, explicit_shadings)
        if not strokes:
            continue
        strokes.sort_in_reading_order()
        table = TableStructure(strokes, **settings).parse(explicit_shadings).to_table_block()
        if isinstance(self._parent, Cell) and table.num_cols * table.num_rows == 1 and table[0][0].bg_color is None:
            continue
        if _straddles(table, table_lines):
            continue  # 글줄이 칸 경계에 걸침 → 표가 아니다
        if table.num_cols == 2 and table.num_rows >= 1 and table[0][0] is not None:
            edge = table[0][0].bbox.x1
            marks = [l.text.strip() for l in table_lines if l.bbox.x1 <= edge + 1]
            if marks and all(len(t) <= 4 for t in marks):
                continue  # 첫 칸이 번호·각주 표시뿐 → 목록·각주이지 표가 아니다
        table.set_stream_table_block()
        tables.append(table)
    self._blocks.assign_to_tables(tables)
    self._shapes.assign_to_tables(tables)


TablesConstructor.stream_tables = _stream_tables

# ---------------------------------------------------------------------------
# 9) 단 나누기: '다음 단' 구역 나누기 대신 단 나누기 문자를 쓴다
# ---------------------------------------------------------------------------
# 원본의 '다음 단' 구역 나누기는 LibreOffice(이 사이트의 Word→PDF 포함)가 무시하고 두 단을 균형 맞춰
# 다시 흘려, 왼쪽 단 글이 오른쪽 단으로 넘어가 섞인다. 단 나누기 문자는 Word·LibreOffice 모두 지킨다.


def _section_make_docx(self, doc):
    set_columns(doc.sections[-1], [c.bbox[2] - c.bbox[0] for c in self], self.space)
    for i, column in enumerate(self):
        if i > 0:
            body = doc.element.body
            last = body[-2] if len(body) > 1 else None  # 맨 끝은 구역 설정(sectPr)
            if last is not None and last.tag == qn('w:p') and doc.paragraphs and doc.paragraphs[-1]._p is last:
                doc.paragraphs[-1].add_run().add_break(WD_BREAK.COLUMN)
            else:
                p = doc.add_paragraph()
                reset_paragraph_format(p, line_spacing=Pt(1))
                p.add_run().add_break(WD_BREAK.COLUMN)
        column.make_docx(doc)


Section.make_docx = _section_make_docx

# ---------------------------------------------------------------------------
# 8) 줄 간격: PDF의 줄 간격을 그대로 (정확히 n pt)
# ---------------------------------------------------------------------------
# 원본은 'n배' 줄 간격을 PDF 글꼴 높이로 계산한다. Word에서는 바꿔 쓴 글꼴(맑은 고딕은 특히 키가 크다)의
# 높이로 다시 곱해져 줄마다 간격이 벌어지고 쪽이 밀린다. 줄 사이 거리를 직접 정하면 글꼴과 무관하게 같다.


def _parse_line_spacing(self):
    for block in self._instances:
        if not block.is_text_block:
            continue
        has_image = any(isinstance(span, ImageSpan) for line in block.lines for span in line.spans)
        if has_image:
            block.parse_relative_line_spacing()  # 글줄 속 그림은 높이가 제각각이라 고정하면 잘린다
        else:
            block.line_space_type = 0
            block.parse_exact_line_spacing()


Blocks._parse_line_spacing = _parse_line_spacing


def _parse_exact_line_spacing(self):
    """원본은 (덩어리 높이 − 첫 줄 높이) ÷ (줄 수 − 1)로 줄 간격을 구하는데, 큰 기호 때문에 두 줄이 한 줄로
    세어지면 간격이 커져 쪽이 넘친다. 줄과 줄 사이 거리의 가운데값을 쓴다."""
    idx = 1 if self.is_horizontal_text else 0
    first = self.lines[0].bbox
    first_h = first[idx + 2] - first[idx]
    block_h = self.bbox[idx + 2] - self.bbox[idx]
    rows = self.lines.group_by_physical_rows()
    count = len(rows)
    if count > 1:
        bottoms = sorted(max(l.bbox[idx + 2] for l in row) for row in rows)
        diffs = sorted(b - a for a, b in zip(bottoms, bottoms[1:]) if b - a > 0.3 * first_h)
        line_space = diffs[len(diffs) // 2] if diffs else (block_h - first_h) / (count - 1)
    else:
        line_space = block_h
    self.line_space = line_space
    self.before_space += first_h - line_space
    if self.before_space < 0:
        self.line_space += self.before_space / max(count, 1)
        self.before_space = 0.0


TextBlock.parse_exact_line_spacing = _parse_exact_line_spacing

# ---------------------------------------------------------------------------
# 7) 그래프·도표(벡터 그림): 축 숫자·범례 같은 짧은 글자까지 함께 그림으로
# ---------------------------------------------------------------------------
# 원본은 벡터 그림 영역을 '글자를 지운' 그림으로 만들고 글자는 따로 남겨, 그래프의 축 숫자들이 그림 밖에
# 흩어지거나 표로 잘못 인식된다. 영역 안 글자가 짧은 이름표뿐이면(그래프) 글자째 그림으로 만든다.
# 문장이 든 영역(둥근 글상자 등)은 글자를 편집할 수 있게 그대로 둔다.


def _is_figure(text_lines):
    texts = [t.strip() for t in text_lines if t.strip()]
    if not texts:
        return True
    long_lines = [t for t in texts if len(t) > 25]
    return sum(len(t) for t in texts) <= 300 and len(long_lines) <= max(1, len(texts) // 6)


# 표 안의 작은 곡선(가운뎃점·글머리 동그라미 등)은 표와 떼어 따로 그림으로 넣는다.
# 원래는 선으로 된 표 묶음에 곡선이 하나라도 섞이면(닫힌 칸 안에 있지 않으면) 표 전체를 그림으로 만들어,
# 표가 사라지고 글자만 흩어진다 (예: 보도자료 표의 '도·소매업' 가운뎃점).
SMALL_CURVE = 6.0  # pt


def _to_shapes_and_images(self, min_svg_gap_dx, min_svg_gap_dy, min_w, min_h, clip_image_res_ratio):
    if self.is_iso_oriented:
        return self.to_shapes(), []
    ie = ImagesExtractor(self.parent.page_engine)
    groups = ie.detect_svg_contours(min_svg_gap_dx, min_svg_gap_dy, min_w, min_h)

    def in_inner(path, contours):
        return any(pymupdf.Rect(b).contains(path.bbox) for b in contours)

    group_paths = [Paths() for _ in groups]
    for path in self._instances:
        for (bbox, inner_bboxes), paths in zip(groups, group_paths):
            if path.bbox.intersects(bbox):
                if not in_inner(path, inner_bboxes):
                    paths.append(path)
                break

    shapes, images = [], []
    clip = lambda rect: ie.clip_page_to_dict(bbox=pymupdf.Rect(rect), rm_image=True, clip_image_res_ratio=clip_image_res_ratio)
    for (bbox, inner_bboxes), paths in zip(groups, group_paths):
        small = [p for p in paths if not p.is_iso_oriented and max(p.bbox.width, p.bbox.height) <= SMALL_CURVE]
        rest = Paths()
        for p in paths:
            if not any(p is q for q in small):
                rest.append(p)
        if len(rest) and rest.is_iso_oriented:  # 표(또는 글자 꾸밈) + 칸 안의 그림 + 작은 곡선
            shapes.extend(rest.to_shapes())
            images.extend(clip(b) for b in inner_bboxes)
            images.extend(clip(pymupdf.Rect(p.bbox) + (-0.5, -0.5, 0.5, 0.5)) for p in small)
        else:  # 그림
            images.append(clip(bbox))
    return shapes, images


def _to_shapes_and_images_patched(self, min_svg_gap_dx=15, min_svg_gap_dy=15, min_w=2, min_h=2, clip_image_res_ratio=3.0):
    shapes, images = _to_shapes_and_images(self, min_svg_gap_dx, min_svg_gap_dy, min_w, min_h, clip_image_res_ratio)
    page = self.parent.page_engine
    regions = getattr(self.parent, '_figure_regions', [])
    text = page.get_text('dict', flags=pymupdf.TEXT_MEDIABOX_CLIP)
    all_lines = [(pymupdf.Rect(l['bbox']), ''.join(sp['text'] for sp in l['spans']))
                 for b in text.get('blocks', []) if b.get('type') == 0 for l in b.get('lines', [])]
    sizes = sorted(sp['size'] for b in text.get('blocks', []) if b.get('type') == 0
                   for l in b.get('lines', []) for sp in l['spans'] if sp['text'].strip())
    body_size = sizes[len(sizes) // 2] if sizes else 10
    label_lines = [(pymupdf.Rect(l['bbox']), ''.join(sp['text'] for sp in l['spans']),
                    max((sp['size'] for sp in l['spans']), default=0))
                   for b in text.get('blocks', []) if b.get('type') == 0 for l in b.get('lines', [])]
    result = []
    for img in images:
        rect = pymupdf.Rect(img['bbox'])
        inside = [t for r, t in all_lines if rect.contains(r) or (r & rect).get_area() > 0.8 * r.get_area()]
        if rect.width > 30 and rect.height > 30 and _is_figure(inside):
            # 틀 바로 바깥의 축 이름·눈금 숫자 같은 짧은 글자도 그림에 넣는다 (본문 글자보다 작은 것만)
            for _ in range(3):
                grown = False
                for r, t, size in label_lines:
                    near = pymupdf.Rect(rect.x0 - 10, rect.y0 - 10, rect.x1 + 10, rect.y1 + 10)
                    if not rect.contains(r) and near.intersects(r) and len(t.strip()) <= 20 and size <= body_size:
                        rect |= r
                        grown = True
                if not grown:
                    break
            pix = page.get_pixmap(clip=rect, matrix=pymupdf.Matrix(clip_image_res_ratio, clip_image_res_ratio))
            img = ImagesExtractor(page)._to_raw_dict(pix, rect)
            regions.append(tuple(rect))
        result.append(img)
    self.parent._figure_regions = regions
    return shapes, result


Paths.to_shapes_and_images = _to_shapes_and_images_patched

_extract_raw_dict = RawPageFitz.extract_raw_dict


def _extract_raw_dict_patched(self, **settings):
    self._figure_regions = []
    raw = _extract_raw_dict(self, **settings)
    regions = [pymupdf.Rect(r) for r in self._figure_regions]
    if regions:
        keep = []
        for block in raw.get('blocks', []):
            if block.get('type', 0) == 0 and 'lines' in block:
                block['lines'] = [l for l in block['lines']
                                  if not any(r.contains(pymupdf.Rect(l['bbox'])) or
                                             (pymupdf.Rect(l['bbox']) & r).get_area() > 0.8 * pymupdf.Rect(l['bbox']).get_area()
                                             for r in regions)]
                if not block['lines']:
                    continue
            keep.append(block)
        raw['blocks'] = keep
    return raw


RawPageFitz.extract_raw_dict = _extract_raw_dict_patched


def drop_margin_stamps(blocks):
    """세로로 쓴 여백 글자(arXiv 표시, 'Downloaded from …' 같은 학술지 다운로드 표시)는 본문 흐름에 넣으면
    단 배치가 무너진다. 본문 글자 영역 바깥의 세로 글줄은 뺀다. (표 안의 세로 글자는 본문 안이라 남는다)"""
    horiz = [l['bbox'] for b in blocks for l in b.get('lines', []) if _horizontal(l)]
    if not horiz:
        return blocks
    x0, x1 = min(b[0] for b in horiz), max(b[2] for b in horiz)
    for block in blocks:
        if 'lines' not in block:
            continue
        block['lines'] = [l for l in block['lines'] if _horizontal(l) or not (
            (l['bbox'][2] <= x0 + 1 or l['bbox'][0] >= x1 - 1) and (l['bbox'][3] - l['bbox'][1]) > 3 * (l['bbox'][2] - l['bbox'][0]))]
    return [b for b in blocks if 'lines' not in b or b['lines']]


# ---------------------------------------------------------------------------
# 6) 2단 문서(논문 등): 페이지 전체에서 단 사이 빈 줄기를 찾아 단을 나눈다
# ---------------------------------------------------------------------------
# 원본은 같은 높이에 놓인 줄끼리 묶어 행마다 단 수를 정하는데, 왼쪽 여백의 세로 글자(arXiv 표시)처럼
# 키 큰 요소가 하나만 있어도 페이지 전체가 '3단 → 1단'으로 처리되어 두 단의 글이 한 단에 뒤섞인다.


def _two_column_sections(raw_page, settings):
    """확실한 2단 페이지면 [1단 띠, 2단 띠, …] 구역 목록을, 아니면 None을 돌려준다."""
    X0, Y0, X1, _ = raw_page.working_bbox
    W = X1 - X0
    lines = [b for b in raw_page.blocks if isinstance(b, Line)]
    body = [l for l in lines if l.is_horizontal_text and len(l.text.strip()) >= 2]
    if len(body) < 12 or W <= 0:
        return None
    total = sum(l.bbox.height for l in body)
    best = None
    x = X0 + 0.3 * W
    while x <= X0 + 0.7 * W:
        cross = sum(l.bbox.height for l in body if l.bbox.x0 < x - 1 < x + 1 < l.bbox.x1)
        if best is None or cross < best[0]:
            best = (cross, x)
        x += 1.0
    cross, gx = best
    left = [l for l in body if l.bbox.x1 <= gx + 1]
    right = [l for l in body if l.bbox.x0 >= gx - 1]
    if cross > 0.25 * total or len(left) < 6 or len(right) < 6:
        return None
    lh, rh = sum(l.bbox.height for l in left), sum(l.bbox.height for l in right)
    if lh < 0.2 * total or rh < 0.2 * total:
        return None
    lw = sorted(l.bbox.width for l in left)[len(left) // 2]
    rw = sorted(l.bbox.width for l in right)[len(right) // 2]
    lcol, rcol = gx - min(l.bbox.x0 for l in left), max(l.bbox.x1 for l in right) - gx
    if not (0.5 <= lcol / rcol <= 2) or lw < 0.6 * lcol or rw < 0.6 * rcol:
        return None  # 단 너비가 너무 다르거나, 한쪽이 짧은 글(양식의 항목 이름 등) 위주 → 2단이 아님

    elements = list(raw_page.blocks) + list(raw_page.shapes)
    side = {}
    spans = []  # 두 단에 걸친 요소들의 세로 범위
    for e in elements:
        b = e.bbox
        if b.x1 <= gx + 2:
            side[id(e)] = 'L'
        elif b.x0 >= gx - 2:
            side[id(e)] = 'R'
        else:
            side[id(e)] = 'F'
            spans.append((b.y0, b.y1))
    spans.sort()
    full = []
    for y0, y1 in spans:
        if full and y0 <= full[-1][1] + 1:
            full[-1][1] = max(full[-1][1], y1)
        else:
            full.append([y0, y1])

    def band_of(e):
        b = e.bbox
        cy = (b.y0 + b.y1) / 2
        for i, (y0, y1) in enumerate(full):
            if side[id(e)] == 'F' and y0 <= b.y0 and b.y1 <= y1:
                return ('F', i)
            if side[id(e)] != 'F' and y0 <= cy <= y1:
                return ('F', i)  # 두 단에 걸친 그림 옆에 끼인 짧은 요소
        before = sum(1 for (y0, _) in full if y0 <= cy)
        return ('C', before)

    bands = {}
    for e in elements:
        bands.setdefault(band_of(e), []).append(e)
    order = sorted(bands, key=lambda k: min(e.bbox.y0 for e in bands[k]))

    sections, y_ref = [], Y0
    for key in order:
        group = bands[key]
        y0, y1 = min(e.bbox.y0 for e in group), max(e.bbox.y1 for e in group)
        cols = [[e for e in group if side[id(e)] == 'L'], [e for e in group if side[id(e)] == 'R']]
        # 각 단에 글줄이 3줄도 안 되는 얇은 띠(저자 이름 한 줄 등)는 단을 나누지 않는다
        few = any(sum(1 for e in col if isinstance(e, Line)) < 3 for col in cols)
        if key[0] == 'F' or few:
            if sections and sections[-1].num_cols == 1:
                column = sections[-1][0]
                column.union_bbox(Collection(group))
                column.add_elements(Collection(group))
            else:
                column = Column().update_bbox((X0, y0, X1, y1))
                column.add_elements(Collection(group))
                section = Section(space=0, columns=[column])
                section.before_space = round(max(0.0, y0 - y_ref), 1)
                sections.append(section)
        else:
            c1 = Column().update_bbox((X0, min(e.bbox.y0 for e in cols[0]), gx, max(e.bbox.y1 for e in cols[0])))
            c1.add_elements(Collection(cols[0]))
            c2 = Column().update_bbox((gx, min(e.bbox.y0 for e in cols[1]), X1, max(e.bbox.y1 for e in cols[1])))
            c2.add_elements(Collection(cols[1]))
            section = Section(space=0, columns=[c1, c2])
            section.before_space = round(max(0.0, y0 - y_ref), 1)
            sections.append(section)
        y_ref = max(y_ref, y1)
    return sections


_parse_section = RawPage.parse_section


def _parse_section_patched(self, **settings):
    try:
        better = _two_column_sections(self, settings)
    except Exception:
        logging.exception('2단 분석 실패: 원래 방식으로 처리')
        better = None
    return better if better is not None else _parse_section(self, **settings)


RawPage.parse_section = _parse_section_patched

# ---------------------------------------------------------------------------
# 12) 표의 행 높이: '정확히' 대신 '최소'로
# ---------------------------------------------------------------------------
# 원본은 행 높이를 '정확히'로 고정한다. Word의 글꼴이 PDF 글꼴보다 조금만 넓어도 칸 안에서 줄이 바뀌는데,
# 고정 높이 때문에 넘친 줄이 가려져 글자가 사라진 것처럼 보인다. '최소'로 두면 행이 늘어날 뿐 다 보인다.

_row_make_docx = Row.make_docx


def _row_make_docx_patched(self, table, idx_row):
    _row_make_docx(self, table, idx_row)
    table.rows[idx_row].height_rule = WD_ROW_HEIGHT.AT_LEAST


Row.make_docx = _row_make_docx_patched

_table_make_docx = TableBlock.make_docx


def _column_widths(block):
    """칸들의 왼쪽·오른쪽 경계에서 열 너비(pt)를 구한다."""
    xs = sorted(x for row in block for cell in row
                if getattr(cell, 'bbox', None) is not None and not cell.bbox.is_empty
                for x in (cell.bbox.x0, cell.bbox.x1))
    bounds = []
    for x in xs:
        if not bounds or x - bounds[-1] > 0.5:
            bounds.append(x)
    return [b - a for a, b in zip(bounds, bounds[1:])]


def _table_make_docx_patched(self, table):
    """표 열 너비를 PDF에서 잰 그대로 고정한다. 자동 맞춤이면 Word·LibreOffice가 내용에 맞춰 열 너비를
    다시 정해, 칸이 몇 개 비어 있는 표에서는 어떤 열이 거의 0이 되어 숫자가 한 글자씩 세로로 늘어선다.
    고정할 때는 표 격자(열 너비)도 같이 적어야 한다. 비워 두면 모든 열이 같은 너비가 되어 넓은 첫 열의 글이 줄바꿈된다."""
    _table_make_docx(self, table)
    table.autofit = False
    widths = _column_widths(self)
    if len(widths) == len(table.columns):
        for column, width in zip(table.columns, widths):
            column.width = Pt(width)


TableBlock.make_docx = _table_make_docx_patched

_cell_make_docx = Cell.make_docx


def _cell_make_docx_patched(self, table, indexes):
    """칸 안 글줄의 앞뒤 빈칸은 뺀다. 이웃 칸 사이의 빈칸이 다음 칸 글 앞에 붙어 오면, 좁은 칸에서 끝 글자가
    다음 줄로 넘어간다 (예: ' (△13.6)')."""
    for block in self.blocks or []:
        lines = getattr(block, 'lines', None) or []
        if not lines:
            continue
        # 문단의 맨 앞과 맨 뒤만 (줄 사이 빈칸은 낱말 사이 띄어쓰기라 남겨야 한다)
        first = [span for span in lines[0].spans if isinstance(span, TextSpan)]
        last = [span for span in lines[-1].spans if isinstance(span, TextSpan)]
        while first and first[0].chars and not first[0].chars[0].c.strip():
            first[0].chars.pop(0)
        while last and last[-1].chars and not last[-1].chars[-1].c.strip():
            last[-1].chars.pop()
    _cell_make_docx(self, table, indexes)


Cell.make_docx = _cell_make_docx_patched

# ---------------------------------------------------------------------------
# 11) 머리글·바닥글: 쪽마다 되풀이되는 위·아래 글줄과 쪽 번호를 Word 머리글·바닥글로
# ---------------------------------------------------------------------------
# 원본은 머리글·바닥글을 본문 글로 넣는다(해당 기능이 비어 있음). 그러면 글꼴 차이로 본문이 조금만 길어져도
# 쪽 번호 한 줄이 다음 쪽으로 밀려 거의 빈 쪽이 생긴다. 진짜 머리글·바닥글로 옮기고 쪽 번호는 자동 번호로 넣는다.

_PAGE_NUM = re.compile(r'^(?P<pre>[-–—(\[]?\s*(?:page|p\.)?\s*)(?P<num>\d{1,4})(?P<post>\s*(?:/\s*\d{1,4}|of\s+\d{1,4})?\s*(?:쪽|페이지|면)?\s*[-–—)\]]?)$', re.I)
_HF = {}


def _norm(text):
    return re.sub(r'\d+', '#', re.sub(r'\s+', '', text))


def _hf_lines(raw_page):
    return [b for b in raw_page.blocks if isinstance(b, Line) and b.is_horizontal_text and b.spans
            and all(isinstance(sp, TextSpan) for sp in b.spans) and b.text.strip()]


def _parse_document(raw_pages):
    _HF.clear()
    if not raw_pages:
        return '', ''
    n = len(raw_pages)
    seen = {}  # (위/아래, 정규화한 글, 반올림한 y) → 나온 쪽 수
    for rp in raw_pages:
        H = rp.height
        for l in _hf_lines(rp):
            zone = 'top' if l.bbox.y1 <= 0.12 * H else 'bottom' if l.bbox.y0 >= 0.88 * H else None
            if zone:
                key = (zone, _norm(l.text), round(l.bbox.y0 / 6))
                seen[key] = seen.get(key, 0) + 1
    for rp in raw_pages:
        H, W = rp.height, rp.width
        lines = _hf_lines(rp)
        if len(lines) < 2:
            continue
        found = {'top': [], 'bottom': []}
        for l in lines:
            zone = 'top' if l.bbox.y1 <= 0.12 * H else 'bottom' if l.bbox.y0 >= 0.88 * H else None
            if not zone:
                continue
            # 같은 높이에 다른 본문 글이 있으면(본문 마지막 줄 등) 머리글·바닥글이 아니다
            if any(o is not l and o.in_same_row(l) for o in lines):
                continue
            text = l.text.strip()
            repeated = n >= 2 and seen.get((zone, _norm(text), round(l.bbox.y0 / 6)), 0) >= max(2, 0.4 * n)
            page_num = bool(_PAGE_NUM.match(text)) and len(text) <= 16
            if repeated or page_num:
                found[zone].append(l)
        # 본문과 겹치지 않게: 머리글은 본문보다 위, 바닥글은 본문보다 아래여야 한다
        body = [l for l in lines if l not in found['top'] and l not in found['bottom']]
        if not body:
            continue
        top_body = min(l.bbox.y0 for l in body)
        bottom_body = max(l.bbox.y1 for l in body)
        found['top'] = [l for l in found['top'] if l.bbox.y1 <= top_body]
        found['bottom'] = [l for l in found['bottom'] if l.bbox.y0 >= bottom_body]
        if not (found['top'] or found['bottom']):
            continue
        info = {}
        for zone in ('top', 'bottom'):
            items = []
            for l in sorted(found[zone], key=lambda l: (l.bbox.y0, l.bbox.x0)):
                cx = (l.bbox.x0 + l.bbox.x1) / 2
                align = 'center' if abs(cx - W / 2) < 0.12 * W else 'right' if cx > W / 2 else 'left'
                span = next((sp for sp in l.spans if isinstance(sp, TextSpan) and sp.text.strip()), None)
                items.append({'text': l.text.strip(), 'align': align, 'size': span.size if span else 10,
                              'font': office_font(span.font, l.text) if span else None,
                              'bold': bool(span and span.flags & 16), 'color': span.color if span else 0})
            if items:
                ys = [l.bbox for l in found[zone]]
                info[zone] = {'items': items, 'y0': min(b.y0 for b in ys), 'y1': max(b.y1 for b in ys)}
        moved = {id(l) for zone in ('top', 'bottom') for l in found[zone]}
        rp.blocks.reset([b for b in rp.blocks if id(b) not in moved])
        rp._hf = info
        _HF[rp.page_engine.number] = info
    return '', ''


Pages._parse_document = staticmethod(_parse_document)

_calculate_margin = RawPage.calculate_margin


def _calculate_margin_patched(self, **settings):
    """본문 영역은 원래대로 넉넉히 두고(바닥글을 뺀 본문 기준), 머리글·바닥글은 그 바깥 여백 안에 들어가게
    위치를 정한다. Word는 본문과 머리글·바닥글이 겹치면 본문을 밀어내므로 겹치지 않게 맞춘다."""
    left, right, top, bottom = _calculate_margin(self, **settings)
    info = getattr(self, '_hf', None)
    if info and self.blocks:
        _, v0, _, v1 = self.blocks.bbox
        H = self.height
        if 'top' in info:
            t = info['top']
            h = t['y1'] - t['y0']
            top = round(min(max(top, h + 4), max(v0 - 1, 1)), 1)
            t['distance'] = max(2.0, min(t['y0'], top - h - 1))
        if 'bottom' in info:
            b = info['bottom']
            h = b['y1'] - b['y0']
            bottom = round(min(max(bottom, h + 4), max(H - v1 - 1, 1)), 1)
            b['distance'] = max(2.0, min(H - b['y1'], bottom - h - 1))
    return left, right, top, bottom


RawPage.calculate_margin = _calculate_margin_patched


def _write_hf(container, items):
    container.is_linked_to_previous = False
    first = True
    for item in items:
        p = container.paragraphs[0] if first else container.add_paragraph()
        first = False
        p.paragraph_format.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
                                        'right': WD_ALIGN_PARAGRAPH.RIGHT}[item['align']]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        m = _PAGE_NUM.match(item['text'])
        parts = [(m.group('pre'), None), (None, 'PAGE'), (m.group('post'), None)] if m else [(item['text'], None)]
        for text, field in parts:
            if field:
                runs = [p.add_run(), p.add_run(), p.add_run('1'), p.add_run()]
                for run, kind in zip((runs[0], runs[3]), ('begin', 'end')):
                    el = OxmlElement('w:fldChar')
                    el.set(qn('w:fldCharType'), kind)
                    run._r.append(el)
                instr = OxmlElement('w:instrText')
                instr.set(qn('xml:space'), 'preserve')
                instr.text = ' PAGE '
                runs[1]._r.append(instr)
                sep = OxmlElement('w:fldChar')
                sep.set(qn('w:fldCharType'), 'separate')
                runs[1]._r.append(sep)
            elif text:
                runs = [p.add_run(text)]
            else:
                continue
            for run in runs:
                run.font.size = Pt(round(item['size'] * 2) / 2)
                run.bold = item['bold'] or None
                if item['font']:
                    run.font.name = item['font']
                    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), item['font'])
                run.font.color.rgb = RGBColor(*rgb_component(item['color']))


_page_make_docx = Page.make_docx


def _page_make_docx_patched(self, doc):
    first = len(doc.sections) if doc.paragraphs else 0
    _page_make_docx(self, doc)
    if not _HF:
        return
    section = doc.sections[first]
    info = _HF.get(self.id, {})
    _write_hf(section.header, info.get('top', {}).get('items', []))
    _write_hf(section.footer, info.get('bottom', {}).get('items', []))
    if 'top' in info:
        section.header_distance = Pt(info['top'].get('distance', info['top']['y0']))
    if 'bottom' in info:
        section.footer_distance = Pt(info['bottom'].get('distance', self.height - info['bottom']['y1']))


Page.make_docx = _page_make_docx_patched

# ---------------------------------------------------------------------------
# 변환
# ---------------------------------------------------------------------------


class _Progress(logging.Handler):
    """pdf2docx의 진행 기록('(3/10) Page 3')을 받아 진행률을 알려 준다."""

    def __init__(self, report):
        super().__init__(logging.INFO)
        self.report = report
        self.stage = 0
        self.failed = []

    def emit(self, record):
        msg = record.getMessage()
        if '[3/4]' in msg:
            self.stage = 3
        elif '[4/4]' in msg:
            self.stage = 4
        m = re.match(r'\((\d+)/(\d+)\) Page', msg)
        if m:
            done, total = int(m.group(1)), int(m.group(2))
            # 페이지 분석이 대부분의 시간이다: 분석 0~85%, 문서 만들기 85~100%
            p = (done - 1) / total * 0.85 if self.stage == 3 else 0.85 + (done - 1) / total * 0.15
            self.report(p)
        m = re.match(r'Ignore page (\d+)', msg)
        if m:
            self.failed.append(int(m.group(1)))


SETTINGS = {'page_margin_factor_bottom': 0.25}


def convert(src, dst, report=lambda p: None, password=None):
    """src PDF를 dst DOCX로 바꾼다. 반환: 변환하지 못하고 빠진 페이지 번호 목록."""
    with pymupdf.open(src) as doc:
        if doc.needs_pass and not (password and doc.authenticate(password)):
            raise RuntimeError('PASSWORD_REQUIRED')
    handler = _Progress(report)
    root = logging.getLogger()
    old_level = root.level
    root.setLevel(logging.INFO)
    root.addHandler(handler)
    try:
        cv = pdf2docx.Converter(src, password=password)
        try:
            cv.convert(dst, multi_processing=False, **SETTINGS)
        finally:
            cv.close()
    finally:
        root.removeHandler(handler)
        root.setLevel(old_level)
    report(1.0)
    return handler.failed
